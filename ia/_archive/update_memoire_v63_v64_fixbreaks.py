# -*- coding: utf-8 -*-
"""
v63 -> v64 : corrige une regression introduite en supprimant les paragraphes
"vides" autour des 3 pages PARTIE : ils contenaient en realite des sauts de
page manuels (<w:br type="page"/>) qui isolaient chaque page partie sur sa
propre page. Reinsere ces sauts de page avant et apres chacun des 3 tableaux
de centrage.
"""
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"C:\Users\DELL\Downloads\MEMOIRE_SI-ENV_v63.docx"
DST = r"C:\Users\DELL\Downloads\MEMOIRE_SI-ENV_v64.docx"

doc = Document(SRC)

def make_pagebreak_paragraph():
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    return p

count = 0
for t in doc.tables:
    txt = t.cell(0, 0).text
    if 'PARTIE' in txt:
        tbl_el = t._tbl
        tbl_el.addprevious(make_pagebreak_paragraph())
        tbl_el.addnext(make_pagebreak_paragraph())
        count += 1
        print(f"[OK] Sauts de page reinseres autour de : {txt.splitlines()[0]}")

assert count == 3, f"Attendu 3 tables PARTIE, trouve {count}"

doc.save(DST)
print(f"\n=== MEMOIRE SAUVEGARDE : {DST} ===")
